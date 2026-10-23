"""
Generador del Documento Oficial Word (.docx) para el Taller 6.
Diseñado con formato institucional de alta calidad para UNIPUTUMAYO.
Stack: React 18, Vite, TypeScript, ConfigCat SDK, Render.
Solución Oficial del Caso Pagos Express, Sprint Review, Retrospective y Liberación Segura.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, title="NOTA IMPORTANTE", border_color="0A2540", bg_color="F0F4F8"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(f"📌 {title}: ")
    run_t.bold = True
    run_t.font.name = "Arial"
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(10, 37, 64)
    
    run_b = p.add_run(text)
    run_b.font.name = "Arial"
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = RGBColor(40, 50, 60)
    doc.add_paragraph()

def style_header(row, col_widths, bg_color="0A2540"):
    for i, cell in enumerate(row.cells):
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        cell.width = Inches(col_widths[i])
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.name = "Arial"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(255, 255, 255)

def style_row(row, col_widths, is_even=False):
    bg_color = "F8FAFC" if is_even else "FFFFFF"
    for i, cell in enumerate(row.cells):
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        cell.width = Inches(col_widths[i])
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(40, 40, 40)

def main():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # ========================== PORTADA ==========================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("INSTITUCIÓN UNIVERSITARIA DEL PUTUMAYO – UNIPUTUMAYO\nFACULTAD DE INGENIERÍA\nTECNOLOGÍA EN DESARROLLO DE SOFTWARE")
    r_inst.bold = True
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(12)
    r_inst.font.color.rgb = RGBColor(10, 37, 64)
    p_inst.paragraph_format.space_after = Pt(35)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("INFORME TÉCNICO - TALLER 6:\nCONTINUOUS DEPLOYMENT, SPRINT REVIEW & RETROSPECTIVE")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(16)
    r_title.font.color.rgb = RGBColor(0, 102, 204)
    p_title.paragraph_format.space_after = Pt(18)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run(
        "Servicio de Feature Flags: ConfigCat\n"
        "Stack Tecnológico: React 18, Vite, TypeScript, ConfigCat SDK, Render\n"
        "Caso Técnico: 'Pagos Express en Liquidación de Trámites Catastrales'\n"
        "Repositorio: https://github.com/franklincordoba2001-ops/taller-6-continuous-deployment"
    )
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(10.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(70, 80, 90)
    p_sub.paragraph_format.space_after = Pt(70)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run(
        "ESTUDIANTE:\nFranklin Sneider Córdoba De La Cruz\n\n"
        "DOCENTE:\nJuan Pablo\n\n"
        "ASIGNATURA / DIPLOMADO:\nIntegración Continua, Continuous Deployment & DevOps\n\n"
        "MOCOA, PUTUMAYO, COLOMBIA\nOCTUBRE DE 2026"
    )
    r_meta.bold = True
    r_meta.font.name = "Arial"
    r_meta.font.size = Pt(10.5)
    r_meta.font.color.rgb = RGBColor(30, 40, 50)
    
    doc.add_page_break()

    # ========================== SECCIÓN 1 ==========================
    doc.add_heading("1. INTRODUCCIÓN Y OBJETIVO", level=1)
    p1 = doc.add_paragraph(
        "El presente informe documenta la resolución técnica y metodológica del caso de estudio de 'Pagos Express' en el marco "
        "del Taller 6: Continuous Deployment, Sprint Review & Retrospective. El objetivo central es adaptar las ceremonias de Scrum "
        "(Sprint Review y Sprint Retrospective) y las políticas de liberación a un modelo maduro de Continuous Deployment (CD) "
        "y Trunk-Based Development (TBD), desacoplando el despliegue técnico a Render de la liberación al usuario final mediante Feature Toggles en ConfigCat."
    )
    p1.paragraph_format.space_after = Pt(12)

    # ========================== SECCIÓN 2 ==========================
    doc.add_heading("2. EN EL SPRINT REVIEW", level=1)
    
    doc.add_heading("2.1 Principios Comparativos: Review Clásico vs. Review en Continuous Deployment", level=2)
    doc.add_paragraph(
        "• Review Clásico: Demostración de un incremento acumulado al final de una iteración de dos semanas; enfoque de 'todo o nada' "
        "habitualmente presentado en local o staging; feedback tardío centrado en 'qué construimos'.\n\n"
        "• Review en CD + TBD: Demostración de valor incremental que ya reside en el entorno de producción real (Render). Las funcionalidades "
        "están desacopladas mediante Feature Toggles (ConfigCat), permitiendo validación empírica y feedback continuo centrado en 'qué valor e impacto recibieron los usuarios'."
    )

    doc.add_heading("2.2 Protocolo de Demostración y Comunicación en el Review (Caso: Pagos Express)", level=2)
    doc.add_paragraph(
        "1. Demo con el toggle en OFF primero (estado real actual):\n"
        "   Se proyecta la aplicación directamente en Render mostrando la experiencia real y actual de los usuarios (formulario estándar de liquidación), "
        "demostrando que el código ya reside en la rama main y en producción sin alterar el comportamiento de los ciudadanos ni generar riesgo alguno.\n\n"
        "2. Activación temporal restringida:\n"
        "   Luego se activa temporalmente solo para el equipo (o un usuario de prueba) y se muestra el flujo de 'Pagar en 1 Clic con Pagos Express', "
        "demostrando la conmutación instantánea del componente sin requerir nuevos commits, compilaciones ni redespliegues.\n\n"
        "3. Gestión rigurosa de expectativas:\n"
        "   Se evita comprometer una fecha de activación al 20% en el Review, condicionando la liberación a métricas y estabilidad demostrada."
    )

    add_callout(
        doc,
        "“La funcionalidad está técnicamente lista y desplegada, pero todavía no la hemos liberado a usuarios reales porque necesitamos "
        "definir métricas de observación y resolver riesgos de estabilidad del pipeline. La campaña de marketing se pausará o se ajustará.”",
        title="DECLARACIÓN TRANSPARENTE DEL PRODUCT OWNER (PO)",
        border_color="0066CC",
        bg_color="F0F7FF"
    )

    # ========================== SECCIÓN 3 ==========================
    doc.add_heading("3. EN LA SPRINT RETROSPECTIVE (ACCIONES CONCRETAS)", level=1)
    
    doc.add_heading("3.1 Nuevos Focos de Inspección Obligatoria en CD", level=2)
    doc.add_paragraph(
        "En Continuous Deployment, la Retrospectiva deja de ser una dinámica abstracta y analiza métricas técnicas de flujo:\n"
        "• Salud del Pipeline: Tasa de builds rotos, tiempos de auto-deploy, estabilidad de health checks y presencia de rollbacks.\n"
        "• Disciplina TBD: Integraciones diarias a main, tiempo de vida de ramas (< 24 horas) e higiene/retiro de feature flags.\n"
        "• Flujo de Valor: Lead Time for Changes y tiempo efectivo en que el usuario recibe valor con observabilidad activa."
    )

    doc.add_heading("3.2 Matriz de Retrospectiva (Estructura Miro)", level=2)
    tbl_retro = doc.add_table(rows=1, cols=3)
    tbl_retro.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w_retro = [2.1, 2.2, 2.2]
    
    r_hdr = tbl_retro.rows[0].cells
    r_hdr[0].text = "🟢 Lo que funcionó bien"
    r_hdr[1].text = "🟡 Lo que generó fricción"
    r_hdr[2].text = "🔴 Lo que rompió el flujo / miedo"
    style_header(tbl_retro.rows[0], col_w_retro, bg_color="0A2540")

    r_row = tbl_retro.add_row()
    r_row.cells[0].text = "• Despliegue automático a Render cumpliendo DoD técnico.\n• Aislamiento efectivo de código sin riesgo al usuario final mediante pagos-express-v1 en ConfigCat.\n• Separación de deploy y release."
    r_row.cells[1].text = "• Tests flaky en CI que demoraron integraciones.\n• Presión comercial del PO / Marketing por activar en viernes sin observabilidad.\n• Solo 4 integraciones en el sprint."
    r_row.cells[2].text = "• Caída de 11 minutos por health check fallido en el último deploy.\n• Rollback manual no automatizado ante fallo de servicio.\n• Acumulación de 3 toggles obsoletos (+45 días)."
    style_row(r_row, col_w_retro)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    doc.add_heading("3.3 Las Tres Acciones Concretas de Mejora Inmediata", level=2)
    doc.add_paragraph(
        "• Acción 1: Arreglar el test flaky en las próximas 24 h (máxima prioridad).\n"
        "  Responsable: Developers. Aislar, refactorizar o mockear de inmediato el test intermitente que bloqueó el CI por 4 días, restableciendo la confianza en las pruebas automáticas.\n\n"
        "• Acción 2: Crear un dashboard mínimo de métricas de pagos (errores + conversión) antes de cualquier activación.\n"
        "  Responsable: Devs + PO. Implementar telemetría en tiempo real sobre errores HTTP 5xx/pasarela y tasa de conversión antes de liberar tráfico externo.\n\n"
        "• Acción 3: Limpiar o eliminar los 3 toggles antiguos en este mismo sprint.\n"
        "  Responsable: Equipo / PO. Erradicar del código base los 3 feature toggles con más de 45 días de vida para pagar la deuda técnica acumulada."
    )

    add_callout(
        doc,
        "“No se activará ningún toggle nuevo a usuarios reales hasta que el health check no haya fallado en los últimos 7 deploys.”",
        title="ACUERDO VINCULANTE DE EQUIPO (RETROSPECTIVA)",
        border_color="D97706",
        bg_color="FFFBEB"
    )

    # ========================== SECCIÓN 4 ==========================
    doc.add_heading("4. ESTRATEGIA DE LIBERACIÓN SEGURA (DECISIÓN FINAL)", level=1)
    
    doc.add_heading("4.1 Evaluación de Decisión: ¿Activar al 20% el Viernes?", level=2)
    doc.add_paragraph(
        "• Decisión Final: NO SE ACTIVA EL 20% ESTE VIERNES.\n"
        "• Justificación Técnica: Incumple las ventanas seguras de despliegue (no liberar antes del fin de semana sin guardia técnica activa); "
        "existe un antecedente reciente de fallo de 11 minutos en health check durante el último deploy; persisten tests flaky en CI; y no existe telemetría previa para monitorear caídas o pagos fallidos."
    )

    doc.add_heading("4.2 Plan Acordado de Liberación Progresiva (5 Pasos)", level=2)
    doc.add_paragraph(
        "1. Paso 1: Resolver test flaky + limpiar toggles viejos (esta semana).\n"
        "2. Paso 2: Implementar métricas básicas de error y conversión (máximo 3 días).\n"
        "3. Paso 3: Activar primero solo para usuarios internos / equipo (1-2 días de observación).\n"
        "4. Paso 4: Si todo está estable → subir al 5% → 20% con observación de 48 h en cada paso.\n"
        "5. Paso 5: Criterio de rollback automático: cualquier aumento > 0,5% en errores de pago activa de inmediato el apagado del toggle en ConfigCat en menos de 60 segundos."
    )

    add_callout(
        doc,
        "“Se protege la confianza del pipeline y la disciplina TBD por encima de la presión comercial de corto plazo. "
        "Se comunica con transparencia al stakeholder y se define un camino claro y seguro de liberación.”",
        title="RESUMEN DE LA DECISIÓN MADURA",
        border_color="10B981",
        bg_color="ECFDF5"
    )

    # ========================== SECCIÓN 5 ==========================
    doc.add_heading("5. REFLEXIÓN PEDAGÓGICA Y CONCLUSIONES", level=1)
    
    add_callout(
        doc,
        "“Este caso es suficientemente difícil porque combina presión de negocio, deuda técnica, fallos recientes de pipeline, "
        "falta de métricas y ruptura de disciplina TBD. Obliga al equipo a negociar, priorizar y aplicar todo lo aprendido en los cuatro talleres.”",
        title="ANÁLISIS DE COMPLEJIDAD Y NEGOCIACIÓN DEL CASO",
        border_color="7C3AED",
        bg_color="F5F3FF"
    )

    doc.add_paragraph(
        "1. Continuous Deployment exige desacoplar el despliegue técnico (automatizado hacia Render) de la liberación de valor (controlada por el PO en ConfigCat).\n"
        "2. El Sprint Review valida hipótesis de valor e impacto real con telemetría, no promesas de fin de ciclo.\n"
        "3. La Sprint Retrospective audita matemáticamente la salud del pipeline y protege las 5 reglas de oro de Trunk-Based Development."
    )

    add_callout(
        doc,
        "“Desplegar a producción ya no es un evento de cierre de sprint, sino una rutina continua e invisible; el Review valida el impacto del valor real y la Retrospectiva protege la salud del pipeline y la disciplina diaria de integración.”",
        title="FRASE X OFICIAL",
        border_color="00D4B2",
        bg_color="F0FFF4"
    )

    output_path = r"c:\Users\Usuario\OneDrive\Escritorio\taller 6\INFORME_TALLER_6_CONTINUOUS_DEPLOYMENT.docx"
    doc.save(output_path)
    print("Documento oficial Word actualizado con éxito en:", output_path)

if __name__ == "__main__":
    main()
